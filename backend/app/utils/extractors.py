from PyPDF2 import PdfReader
import io
import re
from docx import Document
from pptx import Presentation

def clean_format_text(text: str) -> str:
    text = re.sub(r'\n+', ' ', text)
    text = re.sub(r'\s{2,}', ' ', text)
    return text.strip()

def extract_pdf_text(file_bytes: bytes):
    try:
        pdf = PdfReader(io.BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in pdf.pages]
        raw_text =  "\n".join(pages)
        return clean_format_text(raw_text)
    except Exception as e:
        return f"Error reading PDF: {e}"

def extract_document_text(file_bytes: bytes):
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        raw_text = "\n".join(paragraphs)
        return clean_format_text(raw_text)
    except Exception as e:
        return f"Error reading document: {e}"

def extract_pptx_text(file_bytes: bytes):
    try:
        # Ensure we have valid bytes
        if not file_bytes or len(file_bytes) == 0:
            return "Error: Empty file"
            
        # Create a bytes buffer
        file_buffer = io.BytesIO(file_bytes)
        
        # Try to load the presentation
        try:
            prs = Presentation(file_buffer)
        except Exception as e:
            return f"Error loading PowerPoint: {e}"
        
        text_content = []
        
        # Extract text from slides
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    if shape.text and shape.text.strip():
                        text_content.append(shape.text.strip())
                # Also check for text in tables
                elif hasattr(shape, "table"):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text and cell.text.strip():
                                text_content.append(cell.text.strip())
        
        if not text_content:
            return "No text content found in PowerPoint"
            
        raw_text = "\n".join(text_content)
        return clean_format_text(raw_text)
        
    except TypeError as e:
        if "slice indices" in str(e):
            return "Error: Invalid PowerPoint file format or corrupted file"
        return f"Error reading PowerPoint: {e}"
    except Exception as e:
        return f"Error reading PowerPoint: {e}"


    
